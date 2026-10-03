"""Excel KPI report generation and export handlers."""

import io
import json
import logging
from datetime import datetime, timezone, timedelta
from telegram import Update
from telegram.ext import ContextTypes

import config
from models import database as db_app
from views import keyboards as kb
from views.formatting import format_date_ar
from views.messages import _build_dashboard_text
from controllers.common import is_supervisor

logger = logging.getLogger("bot.reports")


async def generate_kpi_excel() -> io.BytesIO:
    """توليد ملف Excel احترافي يحتوي على 3 أوراق عمل مع مخططات بيانية وتنسيق عربي RTL."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.chart import BarChart, Reference
    except ImportError:
        raise ImportError("openpyxl غير مثبت. قم بتشغيل: pip install openpyxl")

    now = datetime.now(timezone(timedelta(hours=3)))
    today_str = now.strftime("%Y-%m-%d")
    month_start = now.replace(day=1).strftime("%Y-%m-%d")
    month_name = kb.AR_MONTHS[now.month - 1] if hasattr(kb, 'AR_MONTHS') else now.strftime('%B')

    # جلب البيانات عبر الاتصال المركزي الموحد
    db = await db_app.get_db()

    # 1. إحصائيات عامة
    async with db.execute("SELECT COUNT(*) FROM requests") as c:
        total_all = (await c.fetchone())[0]
    async with db.execute("SELECT COUNT(*) FROM requests WHERE status='مقبول'") as c:
        total_approved = (await c.fetchone())[0]
    async with db.execute("SELECT COUNT(*) FROM requests WHERE status='مرفوض'") as c:
        total_rejected = (await c.fetchone())[0]
    async with db.execute("SELECT COUNT(*) FROM requests WHERE status='معلق'") as c:
        total_pending = (await c.fetchone())[0]

    # 2. إحصائيات هذا الشهر
    async with db.execute(
        "SELECT COUNT(*) FROM requests WHERE timestamp >= ?", (month_start,)
    ) as c:
        month_total = (await c.fetchone())[0]
    async with db.execute(
        "SELECT COUNT(*) FROM requests WHERE status='مقبول' AND timestamp >= ?", (month_start,)
    ) as c:
        month_approved = (await c.fetchone())[0]
    async with db.execute(
        "SELECT COUNT(*) FROM requests WHERE status='مرفوض' AND timestamp >= ?", (month_start,)
    ) as c:
        month_rejected = (await c.fetchone())[0]

    # 3. أكثر الأقسام طلباً
    async with db.execute(
        "SELECT department, COUNT(*) as cnt FROM requests WHERE department != '' GROUP BY department ORDER BY cnt DESC LIMIT 10"
    ) as c:
        dept_stats = await c.fetchall()

    # 4. توزيع حالات الطلبات حسب القسم
    async with db.execute(
        "SELECT department, status, COUNT(*) as cnt FROM requests WHERE department != '' GROUP BY department, status"
    ) as c:
        dept_status_rows = await c.fetchall()

    # 5. الاتجاه الشهري (آخر 6 أشهر)
    monthly_trend = []
    for i in range(5, -1, -1):
        if now.month - i <= 0:
            m_month = now.month - i + 12
            m_year = now.year - 1
        else:
            m_month = now.month - i
            m_year = now.year
        m_start = f"{m_year}-{m_month:02d}-01"
        if m_month == 12:
            m_end = f"{m_year + 1}-01-01"
        else:
            m_end = f"{m_year}-{m_month + 1:02d}-01"
        async with db.execute(
            "SELECT COUNT(*) FROM requests WHERE timestamp >= ? AND timestamp < ?",
            (m_start, m_end)
        ) as c:
            cnt = (await c.fetchone())[0]
        ar_months = ['يناير','فبراير','مارس','أبريل','مايو','يونيو',
                     'يوليو','أغسطس','سبتمبر','أكتوبر','نوفمبر','ديسمبر']
        monthly_trend.append((ar_months[m_month - 1], cnt))

    # 6. أكثر أنواع التغطية طلباً
    async with db.execute(
        "SELECT coverage_type FROM requests WHERE coverage_type IS NOT NULL AND coverage_type != ''"
    ) as c:
        coverage_rows = await c.fetchall()
    coverage_counter: dict = {}
    for row in coverage_rows:
        coverage_val = row['coverage_type'] or ''
        items = []
        if coverage_val.startswith('['):
            try:
                parsed = json.loads(coverage_val)
                if isinstance(parsed, list):
                    items = parsed
            except Exception:
                pass
        if not items:
            items = [x.strip() for x in coverage_val.split(',') if x.strip()]
        for item in items:
            if item:
                coverage_counter[item] = coverage_counter.get(item, 0) + 1
    coverage_sorted = sorted(coverage_counter.items(), key=lambda x: x[1], reverse=True)[:8]

    # 7. آخر 10 تغطيات مقبولة
    async with db.execute(
        "SELECT event_name, department, date, time FROM requests "
        "WHERE status='مقبول' ORDER BY timestamp DESC LIMIT 10"
    ) as c:
        recent_approved = await c.fetchall()

    # 8. معدل الموافقة
    approval_rate = round((total_approved / total_all * 100), 1) if total_all > 0 else 0
    rejection_rate = round((total_rejected / total_all * 100), 1) if total_all > 0 else 0

    # بناء ملف Excel
    wb = Workbook()

    CLR_HEADER   = "1A3C5E"   # أزرق داكن
    CLR_SUBHEAD  = "2E86C1"   # أزرق متوسط
    CLR_ACCENT   = "E8F4FD"   # أزرق فاتح
    CLR_GREEN    = "27AE60"
    CLR_RED      = "E74C3C"
    CLR_ORANGE   = "F39C12"
    CLR_WHITE    = "FFFFFF"
    CLR_GRAY     = "F2F3F4"
    CLR_GOLD     = "D4AC0D"

    thin = Side(style='thin', color="CCCCCC")
    thick = Side(style='medium', color=CLR_SUBHEAD)
    thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style_header(cell, text, bg=CLR_HEADER, fg=CLR_WHITE, size=12, bold=True):
        cell.value = text
        cell.font = Font(name='Arial', bold=bold, color=fg, size=size)
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border

    def style_cell(cell, text, bg=None, fg="000000", bold=False, align='center'):
        cell.value = text
        cell.font = Font(name='Arial', bold=bold, color=fg, size=10)
        if bg:
            cell.fill = PatternFill("solid", fgColor=bg)
        cell.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
        cell.border = thin_border

    # ══════════════════════════════════════════════════════
    # ورقة 1: ملخص تنفيذي (KPI Overview)
    # ══════════════════════════════════════════════════════
    ws1 = wb.active
    ws1.title = "الملخص التنفيذي"
    ws1.sheet_view.rightToLeft = True
    ws1.column_dimensions['A'].width = 28
    ws1.column_dimensions['B'].width = 18
    ws1.column_dimensions['C'].width = 28
    ws1.column_dimensions['D'].width = 18

    # عنوان رئيسي
    ws1.merge_cells('A1:D1')
    style_header(ws1['A1'], f"📊 تقرير الأداء الإعلامي — {month_name} {now.year}", bg=CLR_HEADER, size=14)
    ws1.row_dimensions[1].height = 35

    ws1.merge_cells('A2:D2')
    style_header(ws1['A2'], f"تاريخ الإصدار: {now.strftime('%Y/%m/%d')} — الساعة {now.strftime('%H:%M')}",
                 bg=CLR_SUBHEAD, size=10, bold=False)
    ws1.row_dimensions[2].height = 22

    ws1.merge_cells('A3:D3')
    style_header(ws1['A3'], "🎯 مؤشرات الأداء الرئيسية (KPIs)", bg=CLR_SUBHEAD, size=11)
    ws1.row_dimensions[3].height = 26

    kpis = [
        ("📋 إجمالي الطلبات (كل الوقت)", total_all, "✅ طلبات مقبولة (كل الوقت)", total_approved),
        ("🗓️ طلبات هذا الشهر", month_total, "✅ مقبولة هذا الشهر", month_approved),
        ("❌ مرفوضة (كل الوقت)", total_rejected, "⏳ معلقة حالياً", total_pending),
        ("📈 معدل الموافقة", f"{approval_rate}%", "📉 معدل الرفض", f"{rejection_rate}%"),
    ]

    for i, (lbl1, val1, lbl2, val2) in enumerate(kpis, start=4):
        ws1.row_dimensions[i].height = 30
        bg = CLR_GRAY if i % 2 == 0 else CLR_ACCENT
        style_cell(ws1.cell(i, 1), lbl1, bg=bg, bold=True, align='right')
        c = ws1.cell(i, 2)
        c.value = val1
        c.font = Font(name='Arial', bold=True, size=13,
                      color=CLR_GREEN if 'مقبول' in str(lbl1) or 'معدل الموافقة' in str(lbl1) else CLR_HEADER)
        c.fill = PatternFill("solid", fgColor=bg)
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = thin_border
        style_cell(ws1.cell(i, 3), lbl2, bg=bg, bold=True, align='right')
        c2 = ws1.cell(i, 4)
        c2.value = val2
        c2.font = Font(name='Arial', bold=True, size=13,
                       color=CLR_RED if 'مرفوض' in str(lbl2) or 'معلق' in str(lbl2) or 'معدل الرفض' in str(lbl2)
                       else CLR_GREEN)
        c2.fill = PatternFill("solid", fgColor=bg)
        c2.alignment = Alignment(horizontal='center', vertical='center')
        c2.border = thin_border

    # فاصل الاتجاه الشهري
    row_offset = len(kpis) + 5
    ws1.merge_cells(f'A{row_offset}:D{row_offset}')
    style_header(ws1[f'A{row_offset}'], "📅 الاتجاه الشهري (آخر 6 أشهر)", bg=CLR_SUBHEAD, size=11)
    ws1.row_dimensions[row_offset].height = 26
    row_offset += 1

    style_header(ws1.cell(row_offset, 1), "الشهر", bg=CLR_HEADER, size=10)
    style_header(ws1.cell(row_offset, 2), "عدد الطلبات", bg=CLR_HEADER, size=10)
    ws1.merge_cells(f'C{row_offset}:D{row_offset}')
    style_header(ws1.cell(row_offset, 3), "ملاحظة", bg=CLR_HEADER, size=10)
    row_offset += 1

    trend_data_start = row_offset
    for m_label, cnt in monthly_trend:
        bg = CLR_GRAY if row_offset % 2 == 0 else CLR_ACCENT
        style_cell(ws1.cell(row_offset, 1), m_label, bg=bg, bold=True)
        c = ws1.cell(row_offset, 2)
        c.value = cnt
        c.font = Font(name='Arial', bold=True, size=11, color=CLR_HEADER)
        c.fill = PatternFill("solid", fgColor=bg)
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = thin_border
        ws1.merge_cells(f'C{row_offset}:D{row_offset}')
        note = "الشهر الحالي" if m_label == month_name else ""
        style_cell(ws1.cell(row_offset, 3), note, bg=bg, fg=CLR_GOLD if note else "888888")
        row_offset += 1
    trend_data_end = row_offset - 1

    chart1 = BarChart()
    chart1.type = "col"
    chart1.grouping = "clustered"
    chart1.title = "الاتجاه الشهري للطلبات"
    chart1.y_axis.title = "عدد الطلبات"
    chart1.style = 10
    chart1.width = 16
    chart1.height = 10
    data_ref = Reference(ws1, min_col=2, min_row=trend_data_start - 1, max_row=trend_data_end)
    cats_ref = Reference(ws1, min_col=1, min_row=trend_data_start, max_row=trend_data_end)
    chart1.add_data(data_ref, titles_from_data=True)
    chart1.set_categories(cats_ref)
    ws1.add_chart(chart1, f"A{row_offset + 1}")

    # ══════════════════════════════════════════════════════
    # ورقة 2: تحليل الأقسام
    # ══════════════════════════════════════════════════════
    ws2 = wb.create_sheet("تحليل الأقسام")
    ws2.sheet_view.rightToLeft = True
    for col, w in zip(['A','B','C','D','E'], [30, 15, 15, 15, 15]):
        ws2.column_dimensions[col].width = w

    ws2.merge_cells('A1:E1')
    style_header(ws2['A1'], "📁 تحليل الطلبات حسب القسم", bg=CLR_HEADER, size=13)
    ws2.row_dimensions[1].height = 32

    headers2 = ["القسم", "إجمالي الطلبات", "مقبولة ✅", "مرفوضة ❌", "معلقة ⏳"]
    for col_i, h in enumerate(headers2, start=1):
        style_header(ws2.cell(2, col_i), h, bg=CLR_SUBHEAD, size=10)
    ws2.row_dimensions[2].height = 26

    dept_map: dict = {}
    for row in dept_status_rows:
        dept = row['department']
        status = row['status']
        cnt = row['cnt']
        if dept not in dept_map:
            dept_map[dept] = {'مقبول': 0, 'مرفوض': 0, 'معلق': 0, 'total': 0}
        dept_map[dept][status] = cnt
        dept_map[dept]['total'] += cnt

    dept_sorted = sorted(dept_map.items(), key=lambda x: x[1]['total'], reverse=True)
    for i, (dept, vals) in enumerate(dept_sorted, start=3):
        bg = CLR_GRAY if i % 2 == 0 else CLR_ACCENT
        ws2.row_dimensions[i].height = 24
        style_cell(ws2.cell(i, 1), dept, bg=bg, bold=True, align='right')
        style_cell(ws2.cell(i, 2), vals['total'], bg=bg, fg=CLR_HEADER, bold=True)
        c_app = ws2.cell(i, 3)
        c_app.value = vals['مقبول']
        c_app.font = Font(name='Arial', bold=True, color=CLR_GREEN, size=10)
        c_app.fill = PatternFill("solid", fgColor=bg)
        c_app.alignment = Alignment(horizontal='center', vertical='center')
        c_app.border = thin_border
        c_rej = ws2.cell(i, 4)
        c_rej.value = vals['مرفوض']
        c_rej.font = Font(name='Arial', bold=True, color=CLR_RED, size=10)
        c_rej.fill = PatternFill("solid", fgColor=bg)
        c_rej.alignment = Alignment(horizontal='center', vertical='center')
        c_rej.border = thin_border
        c_pnd = ws2.cell(i, 5)
        c_pnd.value = vals['معلق']
        c_pnd.font = Font(name='Arial', bold=True, color=CLR_ORANGE, size=10)
        c_pnd.fill = PatternFill("solid", fgColor=bg)
        c_pnd.alignment = Alignment(horizontal='center', vertical='center')
        c_pnd.border = thin_border

    # ══════════════════════════════════════════════════════
    # ورقة 3: أنواع التغطية وآخر المقبولات
    # ══════════════════════════════════════════════════════
    ws3 = wb.create_sheet("التغطيات والأنواع")
    ws3.sheet_view.rightToLeft = True
    for col, w in zip(['A','B','C','D','E'], [30, 14, 26, 20, 16]):
        ws3.column_dimensions[col].width = w

    ws3.merge_cells('A1:B1')
    style_header(ws3['A1'], "🎬 أنواع التغطية الأكثر طلباً", bg=CLR_HEADER, size=12)
    ws3.row_dimensions[1].height = 30

    style_header(ws3.cell(2, 1), "نوع التغطية", bg=CLR_SUBHEAD, size=10)
    style_header(ws3.cell(2, 2), "عدد مرات الطلب", bg=CLR_SUBHEAD, size=10)

    for i, (ctype, cnt) in enumerate(coverage_sorted, start=3):
        bg = CLR_GRAY if i % 2 == 0 else CLR_ACCENT
        ws3.row_dimensions[i].height = 22
        style_cell(ws3.cell(i, 1), ctype, bg=bg, bold=True, align='right')
        c = ws3.cell(i, 2)
        c.value = cnt
        c.font = Font(name='Arial', bold=True, size=11, color=CLR_SUBHEAD)
        c.fill = PatternFill("solid", fgColor=bg)
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = thin_border

    # آخر 10 تغطيات مقبولة
    ws3.merge_cells('C1:E1')
    style_header(ws3['C1'], "✅ آخر 10 تغطيات مقبولة", bg=CLR_HEADER, size=12)
    headers3b = ["اسم الحدث", "القسم", "التاريخ"]
    for col_i, h in enumerate(headers3b, start=3):
        style_header(ws3.cell(2, col_i), h, bg=CLR_SUBHEAD, size=10)

    for i, req in enumerate(recent_approved, start=3):
        bg = CLR_GRAY if i % 2 == 0 else CLR_ACCENT
        ws3.row_dimensions[i].height = 22
        style_cell(ws3.cell(i, 3), req['event_name'], bg=bg, align='right')
        style_cell(ws3.cell(i, 4), req['department'], bg=bg)
        style_cell(ws3.cell(i, 5), format_date_ar(req['date']), bg=bg)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


async def handle_kpi_report(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Handler لزر تقرير الأداء — يولد Excel ويرسله للأدمن."""
    query = update.callback_query
    await query.answer()

    if not await is_supervisor(update.effective_user.id):
        await query.answer("عذراً، هذا الأمر للمشرفين فقط.", show_alert=True)
        return

    await query.edit_message_text(
        "⏳ <b>جاري توليد تقرير الأداء...</b>\nسيصل ملف Excel خلال ثوانٍ.",
        parse_mode="HTML"
    )

    excel_bytes = await generate_kpi_excel()
    now = datetime.now(timezone(timedelta(hours=3)))
    filename = f"KPI_Report_{now.strftime('%Y_%m_%d_%H%M')}.xlsx"

    await context.bot.send_document(
        chat_id=update.effective_chat.id,
        document=excel_bytes,
        filename=filename,
        caption=(
            f"📈 <b>تقرير الأداء الإعلامي</b>\n"
            f"📅 {format_date_ar(now.strftime('%Y-%m-%d'))} — الساعة {now.strftime('%H:%M')}\n\n"
            "يحتوي التقرير على:\n"
            "• ملخص KPIs الرئيسية\n"
            "• تحليل الأقسام\n"
            "• الاتجاه الشهري (آخر 6 أشهر)\n"
            "• أنواع التغطية وآخر المقبولات"
        ),
        parse_mode="HTML"
    )

    # إعادة لوحة التحكم
    try:
        total, pending, approved, rejected = await db_app.get_dashboard_statistics()
        stats_text = _build_dashboard_text(total, pending, approved, rejected)
        await context.bot.send_message(
            chat_id=update.effective_chat.id,
            text=stats_text,
            reply_markup=kb.admin_dashboard_keyboard("مشرف"),
            parse_mode="HTML"
        )
    except Exception as e:
        logger.debug(f"Error resetting admin dashboard after report: {e}")
